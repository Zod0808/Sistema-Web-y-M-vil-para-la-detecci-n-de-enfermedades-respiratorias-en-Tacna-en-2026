"""Convert INFORME_SOLUCION_TECNOLOGICA.md to a formatted .docx per MINEDU spec."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

import sys
_default_src = "INFORME_SOLUCION_TECNOLOGICA.md"
_name = sys.argv[1] if len(sys.argv) > 1 else _default_src
SRC = Path(__file__).parent / _name
DST = Path(__file__).parent / (Path(_name).stem + ".docx")

FONT = "Times New Roman"
FONT_SIZE = Pt(12)
CODE_FONT = "Consolas"
CODE_SIZE = Pt(9)


def set_run_font(run, name=FONT, size=FONT_SIZE, bold=False, italic=False):
    run.font.name = name
    run.font.size = size
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:cs"), name)


def add_inline(paragraph, text, base_bold=False, base_italic=False, code=False):
    """Parse **bold**, *italic*, `code` inside text."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_run_font(run, CODE_FONT if code else FONT,
                         CODE_SIZE if code else FONT_SIZE,
                         bold=base_bold, italic=base_italic)
        token = m.group()
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, FONT, FONT_SIZE, bold=True, italic=base_italic)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, FONT, FONT_SIZE, bold=base_bold, italic=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, CODE_FONT, CODE_SIZE, bold=base_bold, italic=base_italic)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, CODE_FONT if code else FONT,
                     CODE_SIZE if code else FONT_SIZE,
                     bold=base_bold, italic=base_italic)


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run, FONT, FONT_SIZE)
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._element.append(fldChar1)
    run._element.append(instrText)
    run._element.append(fldChar2)


def setup_document(doc):
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_page_number(p)
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = FONT_SIZE


def parse_table(lines, i):
    header = [c.strip() for c in lines[i].strip().strip("|").split("|")]
    i += 2
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
        i += 1
    return header, rows, i


def add_table(doc, header, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Light Grid Accent 1"
    for j, h in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        add_inline(p, h, base_bold=True)
    for r, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            if j >= len(header):
                continue
            cell = table.rows[r].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline(p, val)


def convert():
    md = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup_document(doc)

    i = 0
    while i < len(md):
        line = md[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(md) and not md[i].strip().startswith("```"):
                code_lines.append(md[i])
                i += 1
            i += 1
            for cl in code_lines:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                run = p.add_run(cl if cl else " ")
                set_run_font(run, CODE_FONT, CODE_SIZE)
            doc.add_paragraph()
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sizes = Pt(18)
            elif level == 2:
                sizes = Pt(16)
            elif level == 3:
                sizes = Pt(14)
            else:
                sizes = Pt(12)
            run = p.add_run(text)
            set_run_font(run, FONT, sizes, bold=True)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(md) and re.match(r"^\|[\s:\-|]+\|$", md[i+1].strip()):
            header, rows, i = parse_table(md, i)
            add_table(doc, header, rows)
            doc.add_paragraph()
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            add_inline(p, stripped[2:], base_italic=True)
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, stripped)
        i += 1

    doc.save(DST)
    print(f"OK -> {DST}")


if __name__ == "__main__":
    convert()
